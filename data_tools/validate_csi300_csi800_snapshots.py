"""Validate reconstructed CSI 300 snapshots against all CSMAR CSI 800 weight files."""

from __future__ import annotations

import argparse
import csv
import gzip
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd


CSI800_CODE = "000906"


def norm_code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, int):
        return f"{value:06d}"
    if isinstance(value, float) and value.is_integer():
        return f"{int(value):06d}"
    text = str(value).strip()
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        return str(int(float(text))).zfill(6)
    return text


def norm_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    digits = re.sub(r"\D", "", str(value))
    return f"{digits[:4]}-{digits[4:6]}-{digits[6:8]}" if len(digits) >= 8 else str(value)


def load_reconstructed(daily_gzip: Path, anchor_csv: Path, dates: set[str]) -> dict[str, set[str]]:
    result = {value: set() for value in dates}
    with gzip.open(daily_gzip, "rt", encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            if row["trade_date"] in result:
                result[row["trade_date"]].add(norm_code(row["stock_code"]))
    anchor = pd.read_csv(anchor_csv, dtype={"stock_code": str})
    for snapshot_date, group in anchor.groupby(anchor["snapshot_date"].astype(str)):
        if snapshot_date in result:
            result[snapshot_date] = set(group["stock_code"].map(norm_code))
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--csmar-root", type=Path, required=True)
    parser.add_argument("--daily-gzip", type=Path, required=True)
    parser.add_argument("--anchor-csv", type=Path, required=True)
    parser.add_argument("--dates", nargs="+", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    targets = set(args.dates)
    reconstructed = load_reconstructed(args.daily_gzip, args.anchor_csv, targets)
    found = {value: set() for value in targets}
    file_stats = []
    files = sorted(
        (
            path
            for path in args.csmar_root.rglob("IDX_Smprat*.xlsx")
            if path.parent.name.startswith("\u6307\u6570\u6210\u5206\u80a1\u6743\u91cd\u6587\u4ef6")
        ),
        key=lambda value: str(value).lower(),
    )
    for number, path in enumerate(files, 1):
        print(f"[{number}/{len(files)}] {path.parent.name}/{path.name}", flush=True)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows_read = 0
        csi800_rows = 0
        matched_rows = 0
        min_csi800_date = None
        max_csi800_date = None
        try:
            sheet = workbook[workbook.sheetnames[0]]
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            headers = list(next(rows))
            next(rows, None)
            next(rows, None)
            for raw in rows:
                rows_read += 1
                row = raw[: len(headers)]
                code = norm_code(row[0] if row else None)
                if code and code > CSI800_CODE:
                    break
                if code != CSI800_CODE:
                    continue
                csi800_rows += 1
                row_date = norm_date(row[1])
                if row_date:
                    min_csi800_date = row_date if min_csi800_date is None else min(min_csi800_date, row_date)
                    max_csi800_date = row_date if max_csi800_date is None else max(max_csi800_date, row_date)
                if row_date in found:
                    stock_code = norm_code(row[2])
                    if stock_code:
                        found[row_date].add(stock_code)
                        matched_rows += 1
        finally:
            workbook.close()
        file_stats.append(
            {
                "path": str(path.resolve()),
                "rows_read_until_stop": rows_read,
                "csi800_rows": csi800_rows,
                "matched_rows": matched_rows,
                "min_csi800_date": min_csi800_date,
                "max_csi800_date": max_csi800_date,
            }
        )
    comparisons = {}
    for value in sorted(targets):
        reconstructed_set = reconstructed[value]
        csi800_set = found[value]
        comparisons[value] = {
            "reconstructed_count": len(reconstructed_set),
            "csi800_count": len(csi800_set),
            "not_in_csi800_count": len(reconstructed_set - csi800_set),
            "not_in_csi800_codes": sorted(reconstructed_set - csi800_set),
            "is_subset": bool(reconstructed_set) and not (reconstructed_set - csi800_set),
        }
    result = {
        "single_process": True,
        "single_threaded": True,
        "workbook_count": len(files),
        "target_dates": sorted(targets),
        "files": file_stats,
        "comparisons": comparisons,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"comparisons": comparisons, "files": file_stats}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
