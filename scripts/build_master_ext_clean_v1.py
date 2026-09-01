"""Build MASTER-EXT-clean-v1 from CSMAR XLSX exports.

The pipeline is deliberately single-process and single-threaded.  It preserves
all source workbooks, materializes a cleaned SQLite quote store, computes the
official Qlib Alpha158 factors, the official MASTER 63 market features, and the
official MASTER label::

    Ref($close, -5) / Ref($close, -1) - 1

Raw Alpha158 partitions are stored by instrument.  Final normalized MASTER
inputs are stored as memory-mappable NumPy arrays by calendar year.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import json
import math
import os
import re
import sqlite3
import time
import warnings
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

# Keep every numerical backend on one thread.  These are set before NumPy is
# imported so that BLAS/OpenMP libraries observe them.
for _name in (
    "OMP_NUM_THREADS",
    "OPENBLAS_NUM_THREADS",
    "MKL_NUM_THREADS",
    "VECLIB_MAXIMUM_THREADS",
    "NUMEXPR_NUM_THREADS",
):
    os.environ[_name] = "1"

import numpy as np
import openpyxl
import pandas as pd


START_DATE = "2010-01-04"
END_DATE = "2025-12-31"
FIT_END_DATE = "2020-03-31"
CSI300 = "000300"
CSI500 = "000905"
CSI800 = "000906"
TARGET_INDICES = (CSI300, CSI500, CSI800)
WINDOWS = (5, 10, 20, 30, 60)
EPS = 1e-12
ROBUST_EPS = 1e-12

ALPHA158_NAMES = [
    "KMID", "KLEN", "KMID2", "KUP", "KUP2", "KLOW", "KLOW2", "KSFT", "KSFT2",
    "OPEN0", "HIGH0", "LOW0", "VWAP0",
]
for _op in (
    "ROC", "MA", "STD", "BETA", "RSQR", "RESI", "MAX", "MIN", "QTLU", "QTLD",
    "RANK", "RSV", "IMAX", "IMIN", "IMXD", "CORR", "CORD", "CNTP", "CNTN",
    "CNTD", "SUMP", "SUMN", "SUMD", "VMA", "VSTD", "WVMA", "VSUMP", "VSUMN",
    "VSUMD",
):
    ALPHA158_NAMES.extend(f"{_op}{window}" for window in WINDOWS)
assert len(ALPHA158_NAMES) == 158


def log(message: str) -> None:
    print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {message}", flush=True)


def norm_code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, int):
        return f"{value:06d}"
    if isinstance(value, float) and value.is_integer():
        return f"{int(value):06d}"
    text = str(value).strip()
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        return str(int(float(text))).zfill(6)
    return text


def norm_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    digits = re.sub(r"\D", "", str(value).strip())
    if len(digits) >= 8:
        return f"{digits[:4]}-{digits[4:6]}-{digits[6:8]}"
    return None


def finite(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def qlib_code(code: str) -> str:
    return ("SH" if code.startswith(("5", "6", "9")) else "SZ") + code


def iter_xlsx(path: Path) -> tuple[list[str], Iterable[tuple[Any, ...]], openpyxl.Workbook]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    sheet.reset_dimensions()
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    next(rows, None)
    next(rows, None)
    return headers, rows, workbook


def init_db(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(path)
    con.execute("PRAGMA journal_mode=WAL")
    con.execute("PRAGMA synchronous=NORMAL")
    con.execute("PRAGMA temp_store=MEMORY")
    con.execute("PRAGMA cache_size=-524288")
    con.executescript(
        """
        CREATE TABLE IF NOT EXISTS membership (
            index_code TEXT NOT NULL,
            trade_date TEXT NOT NULL,
            symbol TEXT NOT NULL,
            PRIMARY KEY(index_code, trade_date, symbol)
        ) WITHOUT ROWID;
        CREATE INDEX IF NOT EXISTS membership_symbol_date
            ON membership(index_code, symbol, trade_date);
        CREATE TABLE IF NOT EXISTS quotes (
            trade_date TEXT NOT NULL,
            symbol TEXT NOT NULL,
            open REAL NOT NULL,
            high REAL NOT NULL,
            low REAL NOT NULL,
            close REAL NOT NULL,
            volume REAL,
            amount REAL,
            market_value REAL,
            circulated_market_value REAL,
            source_file TEXT NOT NULL,
            PRIMARY KEY(trade_date, symbol)
        ) WITHOUT ROWID;
        CREATE INDEX IF NOT EXISTS quotes_symbol_date ON quotes(symbol, trade_date);
        CREATE TABLE IF NOT EXISTS factors (
            trade_date TEXT NOT NULL,
            symbol TEXT NOT NULL,
            forward_factor REAL NOT NULL,
            backward_factor REAL NOT NULL,
            cumulative_forward_factor REAL NOT NULL,
            cumulative_backward_factor REAL NOT NULL,
            PRIMARY KEY(trade_date, symbol)
        ) WITHOUT ROWID;
        CREATE INDEX IF NOT EXISTS factors_symbol_date ON factors(symbol, trade_date);
        CREATE TABLE IF NOT EXISTS index_daily (
            index_code TEXT NOT NULL,
            trade_date TEXT NOT NULL,
            close REAL NOT NULL,
            amount REAL NOT NULL,
            PRIMARY KEY(index_code, trade_date)
        ) WITHOUT ROWID;
        CREATE TABLE IF NOT EXISTS quarantine (
            source_file TEXT NOT NULL,
            source_row INTEGER NOT NULL,
            trade_date TEXT,
            symbol TEXT,
            reason TEXT NOT NULL,
            payload TEXT
        );
        """
    )
    return con


def import_csi300_membership(con: sqlite3.Connection, source: Path) -> dict[str, Any]:
    before = con.total_changes
    counts: Counter[str] = Counter()
    with gzip.open(source, "rt", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        batch: list[tuple[str, str, str]] = []
        for row in reader:
            day = row["trade_date"]
            if START_DATE <= day <= END_DATE:
                batch.append((CSI300, day, norm_code(row["stock_code"])))
                counts[day] += 1
                if len(batch) >= 20_000:
                    con.executemany("INSERT OR IGNORE INTO membership VALUES (?,?,?)", batch)
                    batch.clear()
        if batch:
            con.executemany("INSERT OR IGNORE INTO membership VALUES (?,?,?)", batch)
    con.commit()
    bad = {day: count for day, count in counts.items() if count != 300}
    if bad:
        raise ValueError(f"CSI300 membership count violations: {list(bad.items())[:10]}")
    return {
        "inserted": con.total_changes - before,
        "dates": len(counts),
        "count_distribution": dict(Counter(counts.values())),
    }


def scan_index_changes(csmar_root: Path, index_code: str) -> tuple[pd.DataFrame, dict[str, Any]]:
    files = sorted(csmar_root.rglob("IDX_Chgsmp*.xlsx"), key=lambda path: str(path).lower())
    records: list[tuple[str, str, str]] = []
    raw_target = 0
    for number, path in enumerate(files, 1):
        log(f"changes {index_code} {number}/{len(files)} {path.name}")
        headers, rows, workbook = iter_xlsx(path)
        positions = {name: i for i, name in enumerate(headers)}
        try:
            for raw in rows:
                code = norm_code(raw[positions["Indexcd"]])
                if code and code > index_code:
                    break
                if code != index_code:
                    continue
                raw_target += 1
                records.append(
                    (
                        norm_date(raw[positions["Chgsmp01"]]) or "",
                        norm_code(raw[positions["Chgsmp02"]]),
                        str(raw[positions["Chgsmp04"]]).strip(),
                    )
                )
        finally:
            workbook.close()
    frame = pd.DataFrame(records, columns=["change_date", "symbol", "action"])
    frame = frame.drop_duplicates().sort_values(["change_date", "action", "symbol"])
    if frame.empty or set(frame["action"]) - {"1", "2"}:
        raise ValueError(f"Invalid {index_code} change log")
    return frame, {
        "files": len(files),
        "raw_target_rows": raw_target,
        "unique_rows": int(len(frame)),
        "date_min": str(frame.change_date.min()),
        "date_max": str(frame.change_date.max()),
        "actions": {str(k): int(v) for k, v in frame.action.value_counts().items()},
    }


def load_csi800_anchor(csmar_root: Path, anchor_date: str) -> tuple[set[str], dict[str, Any]]:
    files = sorted(csmar_root.rglob("IDX_Smprat*.xlsx"), key=lambda path: str(path).lower())
    codes: set[str] = set()
    scanned: list[str] = []
    for number, path in enumerate(files, 1):
        log(f"CSI800 anchor {number}/{len(files)} {path.name}")
        headers, rows, workbook = iter_xlsx(path)
        positions = {name: i for i, name in enumerate(headers)}
        try:
            for raw in rows:
                code = norm_code(raw[positions["Indexcd"]])
                if code and code > CSI800:
                    break
                if code != CSI800:
                    continue
                day = norm_date(raw[positions["Enddt"]])
                if day == anchor_date:
                    stock = norm_code(raw[positions["Stkcd"]])
                    if stock:
                        codes.add(stock)
        finally:
            workbook.close()
        scanned.append(str(path.resolve()))
        if len(codes) == 800:
            break
    if len(codes) != 800:
        raise ValueError(f"CSI800 anchor {anchor_date} has {len(codes)} codes, expected 800")
    return codes, {"date": anchor_date, "count": len(codes), "files_scanned": scanned}


def index_calendar(con: sqlite3.Connection) -> list[str]:
    rows = con.execute(
        "SELECT trade_date FROM index_daily WHERE index_code=? AND trade_date BETWEEN ? AND ? ORDER BY trade_date",
        (CSI300, START_DATE, END_DATE),
    ).fetchall()
    return [row[0] for row in rows]


def reconstruct_csi800(
    con: sqlite3.Connection,
    csmar_root: Path,
    anchor_date: str,
) -> dict[str, Any]:
    anchor, anchor_audit = load_csi800_anchor(csmar_root, anchor_date)
    changes, changes_audit = scan_index_changes(csmar_root, CSI800)
    grouped: dict[str, dict[str, set[str]]] = defaultdict(lambda: {"adds": set(), "deletes": set()})
    for row in changes.itertuples(index=False):
        grouped[row.change_date]["adds" if row.action == "1" else "deletes"].add(row.symbol)
    calendar = index_calendar(con)
    if not calendar:
        raise ValueError("Index calendar must be extracted before CSI800 reconstruction")
    current = set(anchor)
    reverse_violations: list[dict[str, Any]] = []
    for event_date in sorted(grouped, reverse=True):
        if event_date > anchor_date or event_date <= calendar[0]:
            continue
        adds = grouped[event_date]["adds"]
        deletes = grouped[event_date]["deletes"]
        if adds - current or deletes & current:
            reverse_violations.append(
                {"date": event_date, "missing_adds": sorted(adds - current), "present_deletes": sorted(deletes & current)}
            )
        current.difference_update(adds)
        current.update(deletes)
    if len(current) != 800:
        raise ValueError(f"Reconstructed CSI800 start state has {len(current)} members")
    before = con.total_changes
    forward_violations: list[dict[str, Any]] = []
    event_dates = sorted(day for day in grouped if calendar[0] < day <= calendar[-1])
    event_idx = 0
    counts: Counter[int] = Counter()
    batch: list[tuple[str, str, str]] = []
    for day_index, day in enumerate(calendar):
        if day_index:
            while event_idx < len(event_dates) and event_dates[event_idx] <= day:
                event_date = event_dates[event_idx]
                adds = grouped[event_date]["adds"]
                deletes = grouped[event_date]["deletes"]
                if deletes - current or adds & current:
                    forward_violations.append(
                        {"date": event_date, "missing_deletes": sorted(deletes - current), "present_adds": sorted(adds & current)}
                    )
                current.difference_update(deletes)
                current.update(adds)
                event_idx += 1
        counts[len(current)] += 1
        if not 798 <= len(current) <= 800:
            raise ValueError(f"CSI800 has out-of-range count {len(current)} on {day}")
        batch.extend((CSI800, day, symbol) for symbol in sorted(current))
        if len(batch) >= 40_000:
            con.executemany("INSERT OR IGNORE INTO membership VALUES (?,?,?)", batch)
            batch.clear()
    if batch:
        con.executemany("INSERT OR IGNORE INTO membership VALUES (?,?,?)", batch)
    con.commit()
    return {
        "anchor": anchor_audit,
        "changes": changes_audit,
        "inserted": con.total_changes - before,
        "dates": len(calendar),
        "count_distribution": {str(k): int(v) for k, v in counts.items()},
        "reverse_violations": reverse_violations,
        "forward_violations": forward_violations,
    }



def import_csi800_daily_weights(con: sqlite3.Connection, csmar_root: Path) -> dict[str, Any]:
    """Import authoritative point-in-time CSI800 membership from daily weights."""
    existing = con.execute("SELECT COUNT(*) FROM membership WHERE index_code=?", (CSI800,)).fetchone()[0]
    if existing:
        stats = con.execute("SELECT COUNT(DISTINCT trade_date),MIN(n),MAX(n) FROM (SELECT trade_date,COUNT(*) AS n FROM membership WHERE index_code=? GROUP BY trade_date)", (CSI800,)).fetchone()
        if stats[0] == len(index_calendar(con)) and 798 <= stats[1] <= stats[2] <= 800:
            log(f"reusing complete CSI800 daily membership: rows={existing} dates={stats[0]} counts={stats[1]}-{stats[2]}")
            return {"method": "authoritative daily IDX_Smprat snapshots", "files": 0, "raw_target_rows": 0, "unique_rows": existing, "dates": stats[0], "count_distribution": {}, "inserted": 0, "reused": True}
        con.execute("DELETE FROM membership WHERE index_code=?", (CSI800,))
        con.commit()
    files = sorted(csmar_root.rglob("IDX_Smprat*.xlsx"), key=lambda path: str(path).lower())
    raw_target = 0
    before = con.total_changes
    for number, path in enumerate(files, 1):
        log(f"CSI800 daily weights {number}/{len(files)} {path.name}")
        headers, rows, workbook = iter_xlsx(path)
        pos = {name: i for i, name in enumerate(headers)}
        batch: list[tuple[str, str, str]] = []
        try:
            for raw in rows:
                code = norm_code(raw[pos["Indexcd"]])
                if code and code > CSI800:
                    break
                if code != CSI800:
                    continue
                day = norm_date(raw[pos["Enddt"]])
                if not day or not (START_DATE <= day <= END_DATE):
                    continue
                symbol = norm_code(raw[pos["Stkcd"]])
                if not symbol:
                    continue
                raw_target += 1
                batch.append((CSI800, day, symbol))
                if len(batch) >= 40_000:
                    con.executemany("INSERT OR IGNORE INTO membership VALUES (?,?,?)", batch)
                    batch.clear()
        finally:
            workbook.close()
        if batch:
            con.executemany("INSERT OR IGNORE INTO membership VALUES (?,?,?)", batch)
        con.commit()
    rows = con.execute("SELECT COUNT(*) FROM membership WHERE index_code=?", (CSI800,)).fetchone()[0]
    by_day = con.execute(
        "SELECT trade_date,COUNT(*) FROM membership WHERE index_code=? GROUP BY trade_date ORDER BY trade_date",
        (CSI800,),
    ).fetchall()
    calendar = index_calendar(con)
    dates = [row[0] for row in by_day]
    distribution = Counter(row[1] for row in by_day)
    if dates != calendar:
        missing = sorted(set(calendar) - set(dates))
        extra = sorted(set(dates) - set(calendar))
        raise ValueError(f"CSI800 weight/calendar mismatch missing={missing[:10]} extra={extra[:10]}")
    if not by_day or min(row[1] for row in by_day) < 798 or max(row[1] for row in by_day) > 800:
        raise ValueError(f"CSI800 daily weight counts out of expected range: {dict(distribution)}")
    return {
        "method": "authoritative daily IDX_Smprat snapshots",
        "files": len(files), "raw_target_rows": raw_target, "unique_rows": rows,
        "dates": len(by_day), "count_distribution": {str(k): int(v) for k, v in sorted(distribution.items())},
        "inserted": con.total_changes - before,
    }
def extract_index_daily(con: sqlite3.Connection, csmar_root: Path) -> dict[str, Any]:
    existing = {code: con.execute("SELECT COUNT(*) FROM index_daily WHERE index_code=?", (code,)).fetchone()[0] for code in TARGET_INDICES}
    if len(set(existing.values())) == 1 and min(existing.values()) > 0:
        log(f"reusing complete index_daily table: {existing}")
        return {"files": 0, "raw_rows": {}, "unique_rows": existing, "inserted": 0, "reused": True}
    files = sorted(csmar_root.rglob("IDX_Idxtrd*.xlsx"), key=lambda path: str(path).lower())
    raw_counts: Counter[str] = Counter()
    inserted_before = con.total_changes
    for number, path in enumerate(files, 1):
        log(f"index daily {number}/{len(files)} {path.name}")
        headers, rows, workbook = iter_xlsx(path)
        pos = {name: i for i, name in enumerate(headers)}
        batch: list[tuple[str, str, float, float]] = []
        try:
            for raw in rows:
                code = norm_code(raw[pos["Indexcd"]])
                if code not in TARGET_INDICES:
                    continue
                day = norm_date(raw[pos["Idxtrd01"]])
                close = finite(raw[pos["Idxtrd05"]])
                amount = finite(raw[pos["Idxtrd07"]])
                if day and START_DATE <= day <= END_DATE and close and close > 0 and amount is not None and amount >= 0:
                    raw_counts[code] += 1
                    batch.append((code, day, close, amount))
                    if len(batch) >= 20_000:
                        con.executemany("INSERT OR IGNORE INTO index_daily VALUES (?,?,?,?)", batch)
                        batch.clear()
        finally:
            workbook.close()
        if batch:
            con.executemany("INSERT OR IGNORE INTO index_daily VALUES (?,?,?,?)", batch)
        con.commit()
    unique = {
        code: con.execute("SELECT COUNT(*) FROM index_daily WHERE index_code=?", (code,)).fetchone()[0]
        for code in TARGET_INDICES
    }
    if len(set(unique.values())) != 1 or min(unique.values()) == 0:
        raise ValueError(f"Index daily calendars are inconsistent: {unique}")
    return {"files": len(files), "raw_rows": dict(raw_counts), "unique_rows": unique, "inserted": con.total_changes - inserted_before}


def universe(con: sqlite3.Connection) -> set[str]:
    return {row[0] for row in con.execute("SELECT DISTINCT symbol FROM membership WHERE index_code IN (?,?)", (CSI300, CSI800))}


def extract_quotes(con: sqlite3.Connection, raw_root: Path, symbols: set[str]) -> dict[str, Any]:
    files = sorted(raw_root.rglob("LS_SymbolQuotationD*.xlsx"), key=lambda path: str(path).lower())
    input_rows = kept_rows = invalid_rows = 0
    inserted_before = con.total_changes
    file_audit: list[dict[str, Any]] = []
    for number, path in enumerate(files, 1):
        log(f"quotes {number}/{len(files)} {path.parent.name}/{path.name}")
        headers, rows, workbook = iter_xlsx(path)
        pos = {name: i for i, name in enumerate(headers)}
        batch: list[tuple[Any, ...]] = []
        quarantine: list[tuple[Any, ...]] = []
        local_in = local_keep = local_invalid = 0
        try:
            for source_row, raw in enumerate(rows, 4):
                input_rows += 1
                local_in += 1
                symbol = norm_code(raw[pos["Symbol"]])
                if symbol not in symbols:
                    continue
                day = norm_date(raw[pos["TradingDate"]])
                if not day or not (START_DATE <= day <= END_DATE):
                    continue
                values = [finite(raw[pos[name]]) for name in ("OpenPrice", "HighPrice", "LowPrice", "ClosePrice")]
                reason = None
                if any(value is None or value <= 0 for value in values):
                    reason = "missing_or_nonpositive_ohlc"
                elif values[1] < max(values[0], values[2], values[3]) or values[2] > min(values[0], values[1], values[3]):
                    reason = "ohlc_order_violation"
                volume = finite(raw[pos["Volume"]])
                amount = finite(raw[pos["Amount"]])
                if volume is not None and volume < 0:
                    reason = "negative_volume"
                if amount is not None and amount < 0:
                    reason = "negative_amount"
                if reason:
                    invalid_rows += 1
                    local_invalid += 1
                    quarantine.append((str(path.resolve()), source_row, day, symbol, reason, json.dumps(list(raw), ensure_ascii=False, default=str)))
                else:
                    kept_rows += 1
                    local_keep += 1
                    batch.append(
                        (
                            day, symbol, *values, volume, amount,
                            finite(raw[pos["MarketValue"]]), finite(raw[pos["CirculatedMarketValue"]]),
                            str(path.resolve()),
                        )
                    )
                if len(batch) >= 20_000:
                    con.executemany("INSERT OR IGNORE INTO quotes VALUES (?,?,?,?,?,?,?,?,?,?,?)", batch)
                    batch.clear()
                if len(quarantine) >= 2_000:
                    con.executemany("INSERT INTO quarantine VALUES (?,?,?,?,?,?)", quarantine)
                    quarantine.clear()
        finally:
            workbook.close()
        if batch:
            con.executemany("INSERT OR IGNORE INTO quotes VALUES (?,?,?,?,?,?,?,?,?,?,?)", batch)
        if quarantine:
            con.executemany("INSERT INTO quarantine VALUES (?,?,?,?,?,?)", quarantine)
        con.commit()
        file_audit.append({"path": str(path.resolve()), "input_rows": local_in, "kept_rows": local_keep, "invalid_rows": local_invalid})
    inserted = con.total_changes - inserted_before - invalid_rows
    unique_rows = con.execute("SELECT COUNT(*) FROM quotes").fetchone()[0]
    return {
        "files": len(files), "input_rows": input_rows, "eligible_rows": kept_rows,
        "invalid_rows": invalid_rows, "inserted_changes": inserted, "unique_rows": unique_rows,
        "duplicate_eligible_rows": kept_rows - unique_rows, "file_audit": file_audit,
    }


def extract_factors(con: sqlite3.Connection, raw_root: Path, symbols: set[str]) -> dict[str, Any]:
    files = sorted(raw_root.rglob("TRD_AdjustFactor*.xlsx"), key=lambda path: str(path).lower())
    raw_target = invalid = 0
    before = con.total_changes
    for path in files:
        log(f"factors {path.name}")
        headers, rows, workbook = iter_xlsx(path)
        pos = {name: i for i, name in enumerate(headers)}
        batch: list[tuple[Any, ...]] = []
        try:
            for raw in rows:
                symbol = norm_code(raw[pos["Symbol"]])
                if symbol not in symbols:
                    continue
                day = norm_date(raw[pos["TradingDate"]])
                values = [finite(raw[pos[name]]) for name in ("FwardFactor", "BwardFactor", "CumulateFwardFactor", "CumulateBwardFactor")]
                raw_target += 1
                if not day or any(value is None or value <= 0 for value in values):
                    invalid += 1
                    continue
                batch.append((day, symbol, *values))
                if len(batch) >= 20_000:
                    con.executemany("INSERT OR IGNORE INTO factors VALUES (?,?,?,?,?,?)", batch)
                    batch.clear()
        finally:
            workbook.close()
        if batch:
            con.executemany("INSERT OR IGNORE INTO factors VALUES (?,?,?,?,?,?)", batch)
        con.commit()
    unique_rows = con.execute("SELECT COUNT(*) FROM factors").fetchone()[0]
    reciprocal_bad = con.execute("SELECT COUNT(*) FROM factors WHERE ABS(forward_factor*backward_factor-1)>0.00005").fetchone()[0]
    return {
        "files": len(files), "target_rows": raw_target, "invalid_rows": invalid,
        "unique_rows": unique_rows, "inserted": con.total_changes - before,
        "forward_backward_reciprocal_violations": reciprocal_bad,
    }


def rolling_sum(values: np.ndarray, window: int) -> np.ndarray:
    valid = np.where(np.isfinite(values), values, 0.0)
    prefix = np.concatenate(([0.0], np.cumsum(valid, dtype=np.float64)))
    indices = np.arange(len(values))
    starts = np.maximum(0, indices - window + 1)
    return prefix[indices + 1] - prefix[starts]


def rolling_count(values: np.ndarray, window: int) -> np.ndarray:
    valid = np.isfinite(values).astype(np.int32)
    prefix = np.concatenate(([0], np.cumsum(valid, dtype=np.int32)))
    indices = np.arange(len(values))
    starts = np.maximum(0, indices - window + 1)
    return prefix[indices + 1] - prefix[starts]


def rolling_regression(values: np.ndarray, window: int) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Vectorized equivalent of Qlib's rolling.pyx Slope/Rsquare/Resi."""
    y = values.astype(np.float64, copy=False)
    n_total = len(y)
    idx = np.arange(n_total, dtype=np.float64)
    mask = np.isfinite(y)
    count = rolling_count(y, window).astype(np.float64)
    ysum = rolling_sum(y, window)
    y2sum = rolling_sum(np.where(mask, y * y, np.nan), window)
    isum = rolling_sum(np.where(mask, idx, np.nan), window)
    i2sum = rolling_sum(np.where(mask, idx * idx, np.nan), window)
    iysum = rolling_sum(np.where(mask, idx * y, np.nan), window)
    c = window - idx
    xsum = isum + c * count
    x2sum = i2sum + 2.0 * c * isum + c * c * count
    xysum = iysum + c * ysum
    denom_x = count * x2sum - xsum * xsum
    numer = count * xysum - xsum * ysum
    with np.errstate(divide="ignore", invalid="ignore"):
        slope = numer / denom_x
        denom_y = count * y2sum - ysum * ysum
        rsquare = numer * numer / (denom_x * denom_y)
        intercept = ysum / count - slope * xsum / count
        residual = y - (slope * window + intercept)
    slope[~np.isfinite(slope)] = np.nan
    rsquare[(~np.isfinite(rsquare)) | (rsquare < 0)] = np.nan
    residual[~np.isfinite(residual)] = np.nan
    return slope, rsquare, residual


def rolling_windows(values: np.ndarray, window: int) -> np.ndarray:
    padded = np.concatenate((np.full(window - 1, np.nan), values.astype(np.float64, copy=False)))
    return np.lib.stride_tricks.sliding_window_view(padded, window)


def rolling_rank_current(values: np.ndarray, window: int) -> np.ndarray:
    win = rolling_windows(values, window)
    current = values[:, None]
    valid = np.isfinite(win)
    count = valid.sum(axis=1)
    less = ((win < current) & valid).sum(axis=1)
    equal = ((win == current) & valid).sum(axis=1)
    with np.errstate(divide="ignore", invalid="ignore"):
        result = (less + (equal + 1.0) / 2.0) / count
    result[(~np.isfinite(values)) | (count == 0)] = np.nan
    return result


def rolling_idx(values: np.ndarray, window: int, maximum: bool) -> np.ndarray:
    win = rolling_windows(values, window)
    valid_any = np.isfinite(win).any(axis=1)
    filled = np.where(np.isfinite(win), win, -np.inf if maximum else np.inf)
    result = (np.argmax(filled, axis=1) if maximum else np.argmin(filled, axis=1)).astype(np.float64) + 1.0
    result[~valid_any] = np.nan
    return result


def compute_alpha158(frame: pd.DataFrame) -> np.ndarray:
    o = frame["open"].to_numpy(dtype=np.float64)
    h = frame["high"].to_numpy(dtype=np.float64)
    l = frame["low"].to_numpy(dtype=np.float64)
    c = frame["close"].to_numpy(dtype=np.float64)
    vwap = frame["vwap"].to_numpy(dtype=np.float64)
    volume = frame["volume"].to_numpy(dtype=np.float64)
    series_c = pd.Series(c)
    series_h = pd.Series(h)
    series_l = pd.Series(l)
    series_v = pd.Series(volume)
    with np.errstate(divide="ignore", invalid="ignore"):
        columns: list[np.ndarray] = [
            (c - o) / o,
            (h - l) / o,
            (c - o) / (h - l + EPS),
            (h - np.maximum(o, c)) / o,
            (h - np.maximum(o, c)) / (h - l + EPS),
            (np.minimum(o, c) - l) / o,
            (np.minimum(o, c) - l) / (h - l + EPS),
            (2 * c - h - l) / o,
            (2 * c - h - l) / (h - l + EPS),
            o / c, h / c, l / c, vwap / c,
        ]

    refs_c = np.roll(c, 1)
    refs_c[0] = np.nan
    refs_v = np.roll(volume, 1)
    refs_v[0] = np.nan
    price_delta = c - refs_c
    volume_delta = volume - refs_v
    abs_price_delta = np.abs(price_delta)
    abs_volume_delta = np.abs(volume_delta)
    up_price = np.maximum(price_delta, 0)
    down_price = np.maximum(-price_delta, 0)
    up_volume = np.maximum(volume_delta, 0)
    down_volume = np.maximum(-volume_delta, 0)
    price_ratio = c / refs_c
    log_volume = np.log(volume + 1.0)
    volume_ratio_log = np.log(volume / refs_v + 1.0)
    weighted_change = np.abs(price_ratio - 1.0) * volume

    operators: dict[str, list[np.ndarray]] = defaultdict(list)
    for window in WINDOWS:
        roll_c = series_c.rolling(window, min_periods=1)
        roll_h = series_h.rolling(window, min_periods=1)
        roll_l = series_l.rolling(window, min_periods=1)
        roll_v = series_v.rolling(window, min_periods=1)
        with np.errstate(divide="ignore", invalid="ignore"):
            operators["ROC"].append(series_c.shift(window).to_numpy() / c)
            operators["MA"].append(roll_c.mean().to_numpy() / c)
            operators["STD"].append(roll_c.std().to_numpy() / c)
            slope, rsquare, residual = rolling_regression(c, window)
            operators["BETA"].append(slope / c)
            operators["RSQR"].append(rsquare)
            operators["RESI"].append(residual / c)
            operators["MAX"].append(roll_h.max().to_numpy() / c)
            operators["MIN"].append(roll_l.min().to_numpy() / c)
            operators["QTLU"].append(roll_c.quantile(0.8).to_numpy() / c)
            operators["QTLD"].append(roll_c.quantile(0.2).to_numpy() / c)
            operators["RANK"].append(rolling_rank_current(c, window))
            low_min = roll_l.min().to_numpy()
            high_max = roll_h.max().to_numpy()
            operators["RSV"].append((c - low_min) / (high_max - low_min + EPS))
            idxmax = rolling_idx(h, window, True)
            idxmin = rolling_idx(l, window, False)
            operators["IMAX"].append(idxmax / window)
            operators["IMIN"].append(idxmin / window)
            operators["IMXD"].append((idxmax - idxmin) / window)
            operators["CORR"].append(series_c.rolling(window, min_periods=1).corr(pd.Series(log_volume)).to_numpy())
            operators["CORD"].append(pd.Series(price_ratio).rolling(window, min_periods=1).corr(pd.Series(volume_ratio_log)).to_numpy())
            up_days = rolling_sum(np.where(c > refs_c, 1.0, 0.0), window) / window
            down_days = rolling_sum(np.where(c < refs_c, 1.0, 0.0), window) / window
            # Qlib Mean uses the available prefix length rather than a fixed window.
            prefix_n = np.minimum(np.arange(len(c)) + 1, window)
            up_days *= window / prefix_n
            down_days *= window / prefix_n
            operators["CNTP"].append(up_days)
            operators["CNTN"].append(down_days)
            operators["CNTD"].append(up_days - down_days)
            abs_sum = rolling_sum(abs_price_delta, window)
            gain_sum = rolling_sum(up_price, window)
            loss_sum = rolling_sum(down_price, window)
            operators["SUMP"].append(gain_sum / (abs_sum + EPS))
            operators["SUMN"].append(loss_sum / (abs_sum + EPS))
            operators["SUMD"].append((gain_sum - loss_sum) / (abs_sum + EPS))
            operators["VMA"].append(roll_v.mean().to_numpy() / (volume + EPS))
            operators["VSTD"].append(roll_v.std().to_numpy() / (volume + EPS))
            weighted_s = pd.Series(weighted_change).rolling(window, min_periods=1)
            operators["WVMA"].append(weighted_s.std().to_numpy() / (weighted_s.mean().to_numpy() + EPS))
            vabs_sum = rolling_sum(abs_volume_delta, window)
            vgain_sum = rolling_sum(up_volume, window)
            vloss_sum = rolling_sum(down_volume, window)
            operators["VSUMP"].append(vgain_sum / (vabs_sum + EPS))
            operators["VSUMN"].append(vloss_sum / (vabs_sum + EPS))
            operators["VSUMD"].append((vgain_sum - vloss_sum) / (vabs_sum + EPS))
    for op in (
        "ROC", "MA", "STD", "BETA", "RSQR", "RESI", "MAX", "MIN", "QTLU", "QTLD",
        "RANK", "RSV", "IMAX", "IMIN", "IMXD", "CORR", "CORD", "CNTP", "CNTN",
        "CNTD", "SUMP", "SUMN", "SUMD", "VMA", "VSTD", "WVMA", "VSUMP", "VSUMN", "VSUMD",
    ):
        columns.extend(operators[op])
    result = np.column_stack(columns)
    result[~np.isfinite(result)] = np.nan
    if result.shape[1] != 158:
        raise AssertionError(result.shape)
    return result.astype(np.float32)


def load_adjusted_symbol(con: sqlite3.Connection, symbol: str, calendar: pd.Index) -> tuple[pd.DataFrame, dict[str, Any]]:
    rows = con.execute(
        "SELECT trade_date,open,high,low,close,volume,amount FROM quotes WHERE symbol=? ORDER BY trade_date",
        (symbol,),
    ).fetchall()
    frame = pd.DataFrame(rows, columns=["trade_date", "open", "high", "low", "close", "volume", "amount"])
    if frame.empty:
        frame = pd.DataFrame(index=calendar, columns=["open", "high", "low", "close", "volume", "amount"], dtype=float)
    else:
        frame = frame.set_index("trade_date").reindex(calendar)
    events = con.execute(
        "SELECT trade_date,backward_factor,cumulative_backward_factor FROM factors WHERE symbol=? ORDER BY trade_date",
        (symbol,),
    ).fetchall()
    factor = pd.Series(np.nan, index=calendar, dtype=float)
    if events:
        first_day, first_backward, first_cumulative = events[0]
        events_before = [event for event in events if event[0] <= calendar[0]]
        initial = events_before[-1][2] if events_before else first_cumulative / first_backward
        factor.iloc[0] = initial
        for event_day, _, cumulative in events:
            if event_day in factor.index and event_day > calendar[0]:
                factor.loc[event_day] = cumulative
        factor = factor.ffill()
    else:
        factor[:] = 1.0
    untradable = frame["volume"].isna() | frame["amount"].isna() | frame["volume"].le(0) | frame["amount"].le(0)
    raw_volume = frame["volume"].copy()
    raw_amount = frame["amount"].copy()
    for name in ("open", "high", "low", "close"):
        frame[name] = frame[name] * factor
        frame.loc[untradable, name] = np.nan
    frame["volume"] = raw_volume / factor
    frame.loc[untradable, "volume"] = np.nan
    frame["vwap"] = raw_amount / raw_volume * factor
    frame.loc[untradable, "vwap"] = np.nan
    frame["factor"] = factor
    return frame, {"quote_rows": len(rows), "factor_events": len(events), "untradable_days": int(untradable.sum())}


def membership_dates(con: sqlite3.Connection, index_code: str, symbol: str) -> set[str]:
    return {row[0] for row in con.execute(
        "SELECT trade_date FROM membership WHERE index_code=? AND symbol=? ORDER BY trade_date",
        (index_code, symbol),
    )}


def build_alpha158(con: sqlite3.Connection, output_dir: Path) -> dict[str, Any]:
    raw_dir = output_dir / "alpha158_raw" / "by_instrument"
    raw_dir.mkdir(parents=True, exist_ok=True)
    calendar_values = index_calendar(con)
    calendar = pd.Index(calendar_values, name="trade_date")
    symbols = sorted(universe(con))
    train_chunks: list[np.ndarray] = []
    audit = Counter()
    per_symbol: list[dict[str, Any]] = []
    for number, symbol in enumerate(symbols, 1):
        frame, symbol_audit = load_adjusted_symbol(con, symbol, calendar)
        features = compute_alpha158(frame)
        adjusted_close = frame["close"].to_numpy(dtype=np.float64)
        label = np.full(len(calendar), np.nan, dtype=np.float32)
        label[:-5] = (adjusted_close[5:] / adjusted_close[1:-4] - 1.0).astype(np.float32)
        csi800_dates = membership_dates(con, CSI800, symbol)
        csi300_dates = membership_dates(con, CSI300, symbol)
        selected_800_all = np.fromiter((day in csi800_dates for day in calendar_values), dtype=bool, count=len(calendar_values))
        selected_300_all = np.fromiter((day in csi300_dates for day in calendar_values), dtype=bool, count=len(calendar_values))
        selected = selected_800_all | selected_300_all
        dates = np.array([int(day.replace("-", "")) for day, keep in zip(calendar_values, selected) if keep], dtype=np.int32)
        selected_features = features[selected]
        selected_label = label[selected]
        selected_300 = selected_300_all[selected]
        selected_800 = selected_800_all[selected]
        np.savez_compressed(
            raw_dir / f"{qlib_code(symbol)}.npz",
            dates=dates,
            features=selected_features,
            label=selected_label,
            is_csi300=selected_300,
            is_csi800=selected_800,
        )
        train_mask = dates <= int(FIT_END_DATE.replace("-", ""))
        if train_mask.any():
            train_chunks.append(selected_features[train_mask])
        audit["rows"] += len(dates)
        audit["finite_labels"] += int(np.isfinite(selected_label).sum())
        audit["csi300_rows"] += int(selected_300.sum())
        audit["csi800_rows"] += int(selected_800.sum())
        audit["quote_rows"] += symbol_audit["quote_rows"]
        audit["factor_events"] += symbol_audit["factor_events"]
        per_symbol.append({"symbol": symbol, **symbol_audit, "output_rows": len(dates)})
        if number % 25 == 0 or number == len(symbols):
            log(f"Alpha158 {number}/{len(symbols)} rows={audit['rows']:,}")
    log("computing train-only Alpha158 robust statistics")
    train = np.concatenate(train_chunks, axis=0)
    median = np.nanmedian(train, axis=0).astype(np.float32)
    mad = np.nanmedian(np.abs(train - median), axis=0).astype(np.float32)
    scale = ((mad + ROBUST_EPS) * 1.4826).astype(np.float32)
    del train, train_chunks
    np.savez(output_dir / "alpha158_robust_stats_20100104_20200331.npz", median=median, mad=mad, scale=scale)
    return {
        "symbols": len(symbols), "rows": int(audit["rows"]), "csi300_rows": int(audit["csi300_rows"]), "csi800_rows": int(audit["csi800_rows"]),
        "finite_labels": int(audit["finite_labels"]), "quote_rows": int(audit["quote_rows"]),
        "factor_events": int(audit["factor_events"]), "feature_count": 158,
        "feature_names": ALPHA158_NAMES, "raw_dir": str(raw_dir.resolve()),
        "normalization_fit": [START_DATE, FIT_END_DATE], "per_symbol": per_symbol,
    }


def build_market63(con: sqlite3.Connection, output_dir: Path) -> dict[str, Any]:
    columns: list[str] = []
    all_features: list[np.ndarray] = []
    dates: list[str] | None = None
    for code in TARGET_INDICES:
        rows = con.execute(
            "SELECT trade_date,close,amount FROM index_daily WHERE index_code=? ORDER BY trade_date", (code,)
        ).fetchall()
        frame = pd.DataFrame(rows, columns=["trade_date", "close", "amount"])
        if dates is None:
            dates = frame.trade_date.tolist()
        elif dates != frame.trade_date.tolist():
            raise ValueError("Index calendars differ")
        returns = frame.close / frame.close.shift(1) - 1.0
        per_index = [returns.to_numpy()]
        columns.append(f"{code}_RET1")
        for window in WINDOWS:
            per_index.extend(
                [
                    returns.rolling(window, min_periods=1).mean().to_numpy(),
                    returns.rolling(window, min_periods=1).std().to_numpy(),
                    (frame.amount.rolling(window, min_periods=1).mean() / frame.amount).to_numpy(),
                    (frame.amount.rolling(window, min_periods=1).std() / frame.amount).to_numpy(),
                ]
            )
            columns.extend(
                [f"{code}_RET_MEAN{window}", f"{code}_RET_STD{window}", f"{code}_AMT_MEAN{window}", f"{code}_AMT_STD{window}"]
            )
        all_features.extend(per_index)
    values = np.column_stack(all_features).astype(np.float32)
    if values.shape[1] != 63 or dates is None:
        raise AssertionError(values.shape)
    date_int = np.array([int(day.replace("-", "")) for day in dates], dtype=np.int32)
    train = values[date_int <= int(FIT_END_DATE.replace("-", ""))]
    median = np.nanmedian(train, axis=0).astype(np.float32)
    mad = np.nanmedian(np.abs(train - median), axis=0).astype(np.float32)
    scale = ((mad + ROBUST_EPS) * 1.4826).astype(np.float32)
    normalized = np.clip((values - median) / scale, -3, 3)
    normalized[~np.isfinite(normalized)] = 0.0
    np.savez(
        output_dir / "market63.npz", dates=date_int, raw=values, normalized=normalized.astype(np.float32),
        median=median, mad=mad, scale=scale, names=np.array(columns, dtype="U32"),
    )
    pd.DataFrame(values, index=dates, columns=columns).to_csv(output_dir / "market63_raw.csv.gz", compression="gzip", index_label="trade_date")
    pd.DataFrame(normalized, index=dates, columns=columns).to_csv(output_dir / "market63_normalized.csv.gz", compression="gzip", index_label="trade_date")
    return {
        "dates": len(dates), "date_min": dates[0], "date_max": dates[-1], "feature_count": 63,
        "feature_names": columns, "normalization_fit": [START_DATE, FIT_END_DATE],
    }


def package_years(output_dir: Path) -> dict[str, Any]:
    raw_files = sorted((output_dir / "alpha158_raw" / "by_instrument").glob("*.npz"))
    stats = np.load(output_dir / "alpha158_robust_stats_20100104_20200331.npz")
    median = stats["median"]
    scale = stats["scale"]
    market = np.load(output_dir / "market63.npz")
    market_map = {int(day): row for day, row in zip(market["dates"], market["normalized"])}
    counts: Counter[int] = Counter()
    for path in raw_files:
        with np.load(path) as data:
            counts.update((data["dates"] // 10000).tolist())
    packaged: dict[str, Any] = {}
    package_root = output_dir / "master_input" / "by_year"
    package_root.mkdir(parents=True, exist_ok=True)
    for year in sorted(counts):
        nrows = counts[year]
        year_dir = package_root / f"year={year}"
        year_dir.mkdir(parents=True, exist_ok=True)
        features = np.lib.format.open_memmap(year_dir / "features.npy", mode="w+", dtype=np.float32, shape=(nrows, 221))
        labels = np.lib.format.open_memmap(year_dir / "labels.npy", mode="w+", dtype=np.float32, shape=(nrows,))
        dates_out = np.lib.format.open_memmap(year_dir / "dates.npy", mode="w+", dtype=np.int32, shape=(nrows,))
        instruments = np.lib.format.open_memmap(year_dir / "instruments.npy", mode="w+", dtype="S8", shape=(nrows,))
        is_csi300 = np.lib.format.open_memmap(year_dir / "is_csi300.npy", mode="w+", dtype=np.bool_, shape=(nrows,))
        is_csi800 = np.lib.format.open_memmap(year_dir / "is_csi800.npy", mode="w+", dtype=np.bool_, shape=(nrows,))
        cursor = 0
        for path in raw_files:
            with np.load(path) as data:
                select = data["dates"] // 10000 == year
                count = int(select.sum())
                if not count:
                    continue
                raw = data["features"][select]
                normalized = np.clip((raw - median) / scale, -3, 3)
                normalized[~np.isfinite(normalized)] = 0.0
                selected_dates = data["dates"][select]
                market_values = np.vstack([market_map[int(day)] for day in selected_dates]).astype(np.float32)
                end = cursor + count
                features[cursor:end, :158] = normalized
                features[cursor:end, 158:] = market_values
                labels[cursor:end] = data["label"][select]
                dates_out[cursor:end] = selected_dates
                instruments[cursor:end] = path.stem.encode("ascii")
                is_csi300[cursor:end] = data["is_csi300"][select]
                is_csi800[cursor:end] = data["is_csi800"][select]
                cursor = end
        if cursor != nrows:
            raise AssertionError((year, cursor, nrows))
        features.flush(); labels.flush(); dates_out.flush(); instruments.flush(); is_csi300.flush(); is_csi800.flush()
        del features, labels, dates_out, instruments, is_csi300, is_csi800
        # Stable date/instrument ordering expected by MASTER's daily batching.
        date_view = np.load(year_dir / "dates.npy", mmap_mode="r")
        instrument_view = np.load(year_dir / "instruments.npy", mmap_mode="r")
        order = np.lexsort((np.asarray(instrument_view), np.asarray(date_view)))
        del date_view, instrument_view
        if not np.array_equal(order, np.arange(nrows)):
            for filename in ("features.npy", "labels.npy", "dates.npy", "instruments.npy", "is_csi300.npy", "is_csi800.npy"):
                source = np.load(year_dir / filename, mmap_mode="r")
                temporary = year_dir / (filename + ".sorted")
                target = np.lib.format.open_memmap(temporary, mode="w+", dtype=source.dtype, shape=source.shape)
                chunk = 25_000
                for start in range(0, nrows, chunk):
                    target[start:start + chunk] = source[order[start:start + chunk]]
                target.flush()
                del source, target
                os.replace(temporary, year_dir / filename)
        packaged[str(year)] = {
            "rows": nrows,
            "csi300_rows": int(np.load(year_dir / "is_csi300.npy", mmap_mode="r").sum()),
            "csi800_rows": int(np.load(year_dir / "is_csi800.npy", mmap_mode="r").sum()),
            "finite_labels": int(np.isfinite(np.load(year_dir / "labels.npy", mmap_mode="r")).sum()),
            "feature_shape": [nrows, 221],
            "path": str(year_dir.resolve()),
        }
        log(f"packaged {year}: {nrows:,} rows")
    return {"years": packaged, "feature_count": 221, "factor_count": 158, "market_count": 63}


def export_clean_quotes(con: sqlite3.Connection, output_dir: Path) -> dict[str, Any]:
    target = output_dir / "clean_daily"
    target.mkdir(parents=True, exist_ok=True)
    results: dict[str, Any] = {}
    for year in range(2010, 2026):
        path = target / f"stock_daily_{year}.csv.gz"
        rows = con.execute(
            "SELECT trade_date,symbol,open,high,low,close,volume,amount,market_value,circulated_market_value "
            "FROM quotes WHERE trade_date>=? AND trade_date<=? ORDER BY trade_date,symbol",
            (f"{year}-01-01", f"{year}-12-31"),
        )
        count = 0
        with gzip.open(path, "wt", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["trade_date", "symbol", "open", "high", "low", "close", "volume", "amount", "market_value", "circulated_market_value"])
            for row in rows:
                writer.writerow(row)
                count += 1
        results[str(year)] = {"rows": count, "path": str(path.resolve())}
        log(f"exported clean quotes {year}: {count:,}")
    return results


def validate_outputs(con: sqlite3.Connection, output_dir: Path) -> dict[str, Any]:
    membership = {
        code: {
            "rows": con.execute("SELECT COUNT(*) FROM membership WHERE index_code=?", (code,)).fetchone()[0],
            "dates": con.execute("SELECT COUNT(DISTINCT trade_date) FROM membership WHERE index_code=?", (code,)).fetchone()[0],
            "daily_min": con.execute("SELECT MIN(n) FROM (SELECT COUNT(*) AS n FROM membership WHERE index_code=? GROUP BY trade_date)", (code,)).fetchone()[0],
            "daily_max": con.execute("SELECT MAX(n) FROM (SELECT COUNT(*) AS n FROM membership WHERE index_code=? GROUP BY trade_date)", (code,)).fetchone()[0],
        }
        for code in (CSI300, CSI800)
    }
    year_manifests = sorted((output_dir / "master_input" / "by_year").glob("year=*"))
    checks: dict[str, Any] = {}
    total_rows = 0
    total_300 = 0
    total_800 = 0
    for directory in year_manifests:
        features = np.load(directory / "features.npy", mmap_mode="r")
        labels = np.load(directory / "labels.npy", mmap_mode="r")
        dates = np.load(directory / "dates.npy", mmap_mode="r")
        instruments = np.load(directory / "instruments.npy", mmap_mode="r")
        csi300 = np.load(directory / "is_csi300.npy", mmap_mode="r")
        csi800 = np.load(directory / "is_csi800.npy", mmap_mode="r")
        ordered = bool(np.all((dates[1:] > dates[:-1]) | ((dates[1:] == dates[:-1]) & (instruments[1:] >= instruments[:-1]))))
        checks[directory.name] = {
            "rows": len(dates), "shape": list(features.shape), "ordered": ordered,
            "finite_features": bool(np.isfinite(features).all()), "finite_label_count": int(np.isfinite(labels).sum()),
            "csi300_rows": int(csi300.sum()), "csi800_rows": int(csi800.sum()),
        }
        total_rows += len(dates)
        total_300 += int(csi300.sum())
        total_800 += int(csi800.sum())
    status = "PASS" if (
        membership[CSI300]["rows"] == 300 * membership[CSI300]["dates"]
        and 798 <= membership[CSI800]["daily_min"] <= membership[CSI800]["daily_max"] <= 800
        and total_rows >= max(membership[CSI300]["rows"], membership[CSI800]["rows"])
        and total_800 == membership[CSI800]["rows"]
        and total_300 == membership[CSI300]["rows"]
        and all(item["shape"][1] == 221 and item["ordered"] and item["finite_features"] for item in checks.values())
    ) else "FAIL"
    return {"status": status, "membership": membership, "packaged_rows": total_rows, "packaged_csi300_rows": total_300, "packaged_csi800_rows": total_800, "year_checks": checks}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-root", type=Path, default=Path("datasets/master_extension_raw"))
    parser.add_argument("--csi300-membership", type=Path, default=Path("datasets/csi300_reconstructed_2010_2025/csi300_constituents_daily_2010_2025.csv.gz"))
    parser.add_argument("--output-dir", type=Path, default=Path("datasets/master_ext_clean_v1"))
    parser.add_argument("--audit-dir", type=Path, default=Path("audit_outputs"))
    parser.add_argument("--csi800-anchor-date", default="2025-12-19")
    parser.add_argument("--stage", choices=("all", "extract", "features", "package", "validate"), default="all")
    args = parser.parse_args()
    started = time.time()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.audit_dir.mkdir(parents=True, exist_ok=True)
    con = init_db(args.output_dir / "intermediate" / "master_ext_clean.sqlite")
    audit: dict[str, Any] = {
        "material_passport": {
            "dataset": "MASTER-EXT-clean-v1", "mode": "single-process single-threaded",
            "source_root": str(args.raw_root.resolve()), "start_date": START_DATE, "end_date": END_DATE,
            "alpha158_reference": "microsoft/qlib qlib/contrib/data/loader.py",
            "master_reference": "SJTU-DMTai/MASTER qlib-update/workflow_config_master_Alpha158.yaml",
            "label_formula": "Ref($close,-5)/Ref($close,-1)-1",
        }
    }
    previous_audit_path = args.audit_dir / "master_ext_clean_v1_build_audit.json"
    if args.stage != "all" and previous_audit_path.exists():
        previous_audit = json.loads(previous_audit_path.read_text(encoding="utf-8"))
        previous_audit["material_passport"] = audit["material_passport"]
        audit = previous_audit
    try:
        if args.stage in ("all", "extract"):
            audit["index_daily"] = extract_index_daily(con, args.raw_root / "index_csmar")
            audit["csi300_membership"] = import_csi300_membership(con, args.csi300_membership)
            audit["csi800_membership"] = import_csi800_daily_weights(con, args.raw_root / "index_csmar")
            symbols = universe(con)
            audit["universe_symbols"] = len(symbols)
            audit["quotes"] = extract_quotes(con, args.raw_root, symbols)
            audit["factors"] = extract_factors(con, args.raw_root, symbols)
            audit["clean_quote_exports"] = export_clean_quotes(con, args.output_dir)
        if args.stage in ("all", "features"):
            audit["alpha158"] = build_alpha158(con, args.output_dir)
            audit["market63"] = build_market63(con, args.output_dir)
        if args.stage in ("all", "package"):
            audit["package"] = package_years(args.output_dir)
        if args.stage in ("all", "validate"):
            audit["validation"] = validate_outputs(con, args.output_dir)
        audit["elapsed_seconds"] = round(time.time() - started, 3)
        audit["completed_at"] = datetime.now().isoformat(timespec="seconds")
        path = args.audit_dir / "master_ext_clean_v1_build_audit.json"
        path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
        (args.output_dir / "manifest.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
        log(f"saved audit {path}; elapsed={audit['elapsed_seconds']}s")
    finally:
        con.close()


if __name__ == "__main__":
    warnings.filterwarnings("ignore", category=RuntimeWarning)
    main()







