"""Exhaustively verify target index codes in CSMAR IDX_Smprat workbooks.

Read-only and single-process. Unlike the earlier structural audit, this scan does
not stop when an index code exceeds 000906, so Shenzhen aliases are also checked.
"""

from __future__ import annotations

import argparse
import json
import re
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl


TARGETS = ("000300", "000905", "000906", "399300", "399905", "399906")


def norm_code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, int):
        return f"{value:06d}"
    if isinstance(value, float) and value.is_integer():
        return f"{int(value):06d}"
    text = str(value).strip()
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        return str(int(float(text))).zfill(6)
    return text


def norm_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    digits = re.sub(r"\D", "", text)
    return f"{digits[:4]}-{digits[4:6]}-{digits[6:8]}" if len(digits) >= 8 else text


def blank_stat() -> dict[str, Any]:
    return {
        "rows": 0,
        "min_date": None,
        "max_date": None,
        "dates": set(),
        "folders": set(),
        "files": set(),
        "names": set(),
        "missing_stock_code": 0,
        "missing_weight": 0,
        "rows_by_folder": Counter(),
    }


def update_range(stat: dict[str, Any], value: str | None) -> None:
    if value is None:
        return
    stat["dates"].add(value)
    stat["min_date"] = value if stat["min_date"] is None else min(stat["min_date"], value)
    stat["max_date"] = value if stat["max_date"] is None else max(stat["max_date"], value)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    files = sorted(
        (p for p in args.root.rglob("*.xlsx") if p.parent.name.startswith("指数成分股权重文件")),
        key=lambda p: str(p).lower(),
    )
    stats = {code: blank_stat() for code in TARGETS}
    file_stats: list[dict[str, Any]] = []
    started = time.time()

    for number, path in enumerate(files, 1):
        print(f"[{number}/{len(files)}] {path.parent.name}/{path.name}", flush=True)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows_read = 0
        target_rows = Counter()
        distinct_codes = set()
        try:
            sheet = workbook[workbook.sheetnames[0]]
            declared_dimension = sheet.calculate_dimension()
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            headers = list(next(rows))
            next(rows, None)
            next(rows, None)
            width = len(headers)
            for raw in rows:
                row = raw[:width]
                rows_read += 1
                if not row:
                    continue
                code = norm_code(row[0])
                if code:
                    distinct_codes.add(code)
                if code not in stats:
                    continue
                target_rows[code] += 1
                stat = stats[code]
                stat["rows"] += 1
                stat["folders"].add(path.parent.name)
                stat["files"].add(str(path.resolve()))
                stat["rows_by_folder"][path.parent.name] += 1
                end_date = norm_date(row[1] if len(row) > 1 else None)
                update_range(stat, end_date)
                stock_code = norm_code(row[2] if len(row) > 2 else None)
                if not stock_code:
                    stat["missing_stock_code"] += 1
                if len(row) > 3 and row[3] is not None:
                    stat["names"].add(str(row[3]))
                if len(row) <= 4 or row[4] is None or row[4] == "":
                    stat["missing_weight"] += 1
        finally:
            workbook.close()
        file_stats.append(
            {
                "path": str(path.resolve()),
                "declared_dimension": declared_dimension,
                "rows_read": rows_read,
                "distinct_index_codes": len(distinct_codes),
                "min_index_code": min(distinct_codes) if distinct_codes else None,
                "max_index_code": max(distinct_codes) if distinct_codes else None,
                "target_rows": dict(target_rows),
            }
        )

    output_stats: dict[str, Any] = {}
    for code, stat in stats.items():
        output_stats[code] = {
            **{k: v for k, v in stat.items() if k not in {"dates", "folders", "files", "names", "rows_by_folder"}},
            "unique_dates": len(stat["dates"]),
            "folders": sorted(stat["folders"]),
            "files": sorted(stat["files"]),
            "constituent_name_count": len(stat["names"]),
            "rows_by_folder": dict(stat["rows_by_folder"]),
        }

    result = {
        "root": str(args.root.resolve()),
        "single_process": True,
        "exhaustive_no_early_stop": True,
        "targets": TARGETS,
        "workbook_count": len(files),
        "workbook_bytes": sum(p.stat().st_size for p in files),
        "stats": output_stats,
        "files": file_stats,
        "elapsed_seconds": round(time.time() - started, 3),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(output_stats, ensure_ascii=True, indent=2), flush=True)
    print(f"Wrote {args.output}; elapsed={result['elapsed_seconds']}s", flush=True)


if __name__ == "__main__":
    main()
